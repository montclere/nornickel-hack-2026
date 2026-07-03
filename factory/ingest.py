# -*- coding: utf-8 -*-
"""Универсальный приём разнородных материалов → фрагменты с провенансом.

Человек кладёт ЧТО УГОДНО (PDF/DOCX/TXT/XLSX/CSV) — здесь всё автоматически
разбирается в единый список Chunk. Структуру НЕ хардкодим: текст идёт как текст,
таблицы — как строки с координатами. Дальнейшее извлечение сущностей — в extract.py.

Chunk = {text, source, locator, kind: 'prose'|'table', meta}. meta несёт доступные
метаданные (файл, страница/лист, дата файла).
"""
from __future__ import annotations

import csv
import glob
import html
import os
import re
import zipfile
from dataclasses import dataclass, field

from factory.config import MAX_CHUNK_CHARS, OCR_DPI, OCR_ENABLED, OCR_MIN_CHARS, OCR_PDF_MAX_PAGES

# картинки, которые распознаём через OCR (mime для Yandex Vision)
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
              ".tif": "image/tiff", ".tiff": "image/tiff", ".bmp": "image/bmp",
              ".webp": "image/webp"}
IMAGE_EXTS = set(IMAGE_MIME)


@dataclass
class Chunk:
    text: str
    source: str
    locator: str
    kind: str = "prose"
    meta: dict = field(default_factory=dict)


def _default_ocr():
    """OCR-клиент Yandex, если включён и готов; иначе None (картинки → needs_ocr)."""
    if not OCR_ENABLED:
        return None
    try:
        from factory.ocr import YandexOCR
        o = YandexOCR()
        return o if o.ready else None
    except Exception:  # noqa: BLE001
        return None


def ingest(paths, ocr="auto") -> list:
    """paths — файл, папка или список. Возвращает список Chunk по всем поддерж. файлам.

    ocr: "auto" — поднять Yandex OCR, если доступен; None — без OCR; или готовый клиент."""
    if isinstance(paths, str):
        paths = [paths]
    files = []
    for p in paths:
        if os.path.isdir(p):
            files += glob.glob(os.path.join(p, "**", "*"), recursive=True)
        else:
            files.append(p)
    if ocr == "auto":
        ocr = _default_ocr()
    out = []
    for f in sorted(set(files)):
        if os.path.isfile(f):
            out += ingest_file(f, ocr)
    return out


def ingest_file(path: str, ocr=None) -> list:
    ext = os.path.splitext(path)[1].lower()
    name = os.path.basename(path)
    meta = {"file": name, "mtime": _mtime(path)}
    try:
        if ext == ".pdf":
            return _pdf(path, name, meta, ocr)
        if ext in IMAGE_EXTS:
            return _image(path, name, meta, ocr)
        if ext == ".docx":
            return _docx(path, name, meta)
        if ext in (".txt", ".md"):
            return [Chunk(open(path, encoding="utf-8", errors="ignore").read(),
                          name, name, "prose", meta)]
        if ext == ".xlsx":
            return _xlsx(path, name, meta)
        if ext == ".csv":
            return _csv(path, name, meta)
    except Exception as e:  # noqa: BLE001
        return [Chunk("", name, name, "prose", {**meta, "error": str(e)})]
    return []


def _mtime(path):
    try:
        import datetime
        return datetime.date.fromtimestamp(os.path.getmtime(path)).isoformat()
    except Exception:  # noqa: BLE001
        return None


def _pdf(path, name, meta, ocr=None):
    import fitz
    doc = fitz.open(path)
    out, ocr_pages = [], 0
    for i in range(doc.page_count):
        page = doc[i]
        t = page.get_text().strip()
        if len(t) >= 40:
            out.append(Chunk(t, name, f"{name}:стр.{i+1}", "prose",
                             {**meta, "page": i + 1}))
        elif ocr is not None and ocr.ready and ocr_pages < OCR_PDF_MAX_PAGES:
            # страница без текстового слоя (скан) → растеризуем и распознаём
            ocr_pages += 1
            try:
                ot = ocr.recognize(page.get_pixmap(dpi=OCR_DPI).tobytes("png"), "image/png")
            except Exception:  # noqa: BLE001
                ot = ""
            if len(ot.strip()) >= OCR_MIN_CHARS:
                out.append(Chunk(ot, name, f"{name}:стр.{i+1}", "prose",
                                 {**meta, "page": i + 1, "ocr": True}))
    if not out:
        out.append(Chunk("", name, name, "prose", {**meta, "needs_ocr": True}))
    return out


def _image(path, name, meta, ocr=None):
    """Картинка (скан регламента, схема флотации) → текст через Yandex Vision OCR."""
    meta = {**meta, "ocr": True}
    if ocr is None or not ocr.ready:
        return [Chunk("", name, name, "prose", {**meta, "needs_ocr": True})]
    mime = IMAGE_MIME.get(os.path.splitext(path)[1].lower(), "image/png")
    text = ocr.recognize(open(path, "rb").read(), mime)
    tail = {} if len(text.strip()) >= OCR_MIN_CHARS else {"needs_ocr": True}
    return [Chunk(text, name, name, "prose", {**meta, **tail})]


def _docx(path, name, meta):
    xml = zipfile.ZipFile(path).read("word/document.xml").decode("utf-8", "ignore")
    text = html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"</w:p>", "\n", xml)))
    paras = [p.strip() for p in text.splitlines() if len(p.strip()) > 3]
    return [Chunk("\n".join(paras), name, name, "prose", meta)]


def _xlsx(path, name, meta):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for r, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = [str(c) for c in row if c not in (None, "")]
            if vals:
                rows.append(f"[{ws.title}!{r}] " + " | ".join(vals))
        if rows:
            out.append(Chunk("\n".join(rows), name, f"{name}:{ws.title}", "table",
                             {**meta, "sheet": ws.title}))
    return out


def _csv(path, name, meta):
    rows = []
    with open(path, encoding="utf-8", errors="ignore", newline="") as f:
        for i, row in enumerate(csv.reader(f), 1):
            if any(row):
                rows.append(f"[{i}] " + " | ".join(row))
    return [Chunk("\n".join(rows), name, name, "table", meta)] if rows else []


def split(chunks, max_chars=MAX_CHUNK_CHARS) -> list:
    """Резать длинные прозаические фрагменты по абзацам (окно LLM)."""
    out = []
    for c in chunks:
        if c.kind != "prose" or len(c.text) <= max_chars:
            out.append(c); continue
        buf, size = [], 0
        for para in c.text.split("\n"):
            if size + len(para) > max_chars and buf:
                out.append(Chunk("\n".join(buf), c.source, c.locator, c.kind, c.meta))
                buf, size = [], 0
            buf.append(para); size += len(para) + 1
        if buf:
            out.append(Chunk("\n".join(buf), c.source, c.locator, c.kind, c.meta))
    return out
