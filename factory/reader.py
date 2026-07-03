# -*- coding: utf-8 -*-
"""Чтение отчёта по хвостам → структурированный профиль потерь. ДЕТЕРМИНИРОВАННО.

Никакого LLM. Только парсинг Excel по стабильным подписям-якорям (openpyxl), заданным
СХЕМОЙ (см. schema.py) — не хардкодом. Смена формата отчёта = другая схема, этот файл
не меняется. Один и тот же файл + схема → один и тот же профиль (побайтово).

Классы крупности НЕ фильтруются по белому списку — распознаются и упорядочиваются
динамически, по числу в самой подписи (`class_sort_key`). Так парсер не роняет
классы, которых не было в примерах, если у другой фабрики иная разбивка по крупности.
"""
from __future__ import annotations

import os
import re
from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from factory.schema import DEFAULT_SCHEMA, ReportSchema

# --- обратная совместимость: остальные модули (generator/metrics/analysis/intent)
# читают ЭТИ константы как «текущий известный домен» (Cu-Ni хвосты). Они выводятся из
# DEFAULT_SCHEMA — единственного источника истины, а не задаются заново.
ELEMENT_SYMBOLS = DEFAULT_SCHEMA.element_symbols()
PRIMARY_ELEMENT = DEFAULT_SCHEMA.primary_element()
LIBERATED = DEFAULT_SCHEMA.liberated_form
LOCKED = DEFAULT_SCHEMA.locked_form


def _num(v):
    return v if isinstance(v, (int, float)) else None


def _norm_class(label: str) -> str:
    """«-71 + 45 мкм», « -71 +45» → «-71+45» (канон, без пробелов).

    Подписи класса в таблице и в блоках минералогии расходятся по пробелам —
    убираем все пробелы, чтобы одно и то же читалось одинаково."""
    s = re.sub(r"мкм", "", str(label))
    return re.sub(r"\s+", "", s)


def class_sort_key(label: str):
    """Порядок крупный→тонкий БЕЗ белого списка — по числу(ам) в подписи класса.
    «+125» (>125, открыт сверху) считается чуть крупнее «-125+71» (та же верхняя
    граница 125, но диапазон замкнут). Нераспознанное — в конец, не роняет парсер."""
    nums = [int(x) for x in re.findall(r"\d+", label)]
    if not nums:
        return (10 ** 9, 1)
    upper = max(nums)
    is_open_plus = label.strip().startswith("+")
    return (-upper, 0 if is_open_plus else 1)


def class_upper_micron(label: str) -> float | None:
    """Верхняя граница класса в мкм (для порога «тонкий класс» в rules.py)."""
    nums = [int(x) for x in re.findall(r"\d+", label)]
    return float(max(nums)) if nums else None


@dataclass
class FormLoss:
    """Потеря элемента в конкретной минеральной форме внутри класса крупности."""
    form: str
    element: str
    tonnes: float
    pct: float
    recoverable: bool
    cell: str


@dataclass
class ClassLoss:
    """Потери по классу крупности."""
    size_class: str
    tonnes: dict = field(default_factory=dict)      # {"Ni": т, "Cu": т}
    cells: dict = field(default_factory=dict)        # провенанс {"Ni": "E200"}
    forms: list = field(default_factory=list)        # list[FormLoss]

    def recoverable_tonnes(self, element: str) -> float:
        return round(sum(f.tonnes for f in self.forms
                         if f.element == element and f.recoverable), 1)

    def dominant_recoverable_form(self, element: str) -> str | None:
        cand = [f for f in self.forms if f.element == element and f.recoverable]
        return max(cand, key=lambda f: f.tonnes).form if cand else None


@dataclass
class TailingsProfile:
    fabric: str
    elements: dict = field(default_factory=dict)     # {"Ni": {"grade": .., "cell": ..}}
    classes: list = field(default_factory=list)      # list[ClassLoss], в порядке крупности
    source: str = ""
    schema_name: str = ""
    warnings: list = field(default_factory=list)     # сигналы «парс мог сбиться»

    def by_class(self, sc: str) -> ClassLoss | None:
        return next((c for c in self.classes if c.size_class == sc), None)


class TailingsReader:
    """Парсер отчёта по хвостам. Секция «отвальные хвосты общие» (итоговые потери).

    `schema` — ReportSchema (см. schema.py). По умолчанию DEFAULT_SCHEMA (известный
    формат). Другой формат отчёта → своя схема (файл или объект), код не меняется."""

    def __init__(self, path: str, schema: ReportSchema = DEFAULT_SCHEMA):
        self.path = path
        self.schema = schema
        self.fabric = re.sub(r"Хвосты\s*|\.xlsx|_\d+", "", os.path.basename(path)).strip()

    def _element_columns(self, grid, header_row):
        """Найти колонки (%, т) каждого элемента по подписям возле шапки (label или
        symbol из схемы). Ищем в строках header±2; подпись стоит над парой (%, т).
        Не нашли все элементы схемы — используем документированную раскладку по
        умолчанию (C=доля, затем пары %,т по каждому элементу схемы подряд)."""
        sc = self.schema
        default_cols = {}
        col = sc.anchor_col + 2  # после ANCHOR_COL и колонки "доля класса"
        for e in sc.elements:
            default_cols[e.symbol] = (col, col + 1)
            col += 2

        found = {}
        for (r, c), v in grid.items():
            if abs(r - header_row) > 2 or c <= sc.anchor_col:
                continue
            s = str(v).strip()
            el = next((e.symbol for e in sc.elements if e.label in s or e.symbol == s), None)
            if el and el not in found:
                found[el] = (c, c + 1)
        if set(found) == set(default_cols) and len({c for c, _ in found.values()}) == len(found):
            return found
        return default_cols

    def read(self) -> TailingsProfile:
        sc = self.schema
        ws = load_workbook(self.path, data_only=True).active
        grid = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row
                if c.value is not None}
        prof = TailingsProfile(fabric=self.fabric, source=os.path.basename(self.path),
                               schema_name=sc.name)

        # --- секция: берём ПОСЛЕДНЮЮ таблицу с якорем схемы (отвальные общие) ---
        hdr_rows = sorted(r for (r, col), v in grid.items()
                          if col == sc.anchor_col and sc.size_table_anchor in str(v))
        if not hdr_rows:
            prof.warnings.append(
                f"не найден якорь «{sc.size_table_anchor}» — формат не распознан "
                f"(другой отчёт? нужна другая схема? см. schema.py/schema_bootstrap.py). "
                f"Профиль пуст.")
            return prof
        start = hdr_rows[-1]

        ecol = self._element_columns(grid, start)  # {элемент: (pct_col, tonnes_col)}

        # --- таблица классов крупности ---
        r = start + 1
        classes: dict[str, ClassLoss] = {}
        while r < start + 12:
            b = grid.get((r, sc.anchor_col))
            if b is None:
                r += 1; continue
            if str(b).startswith(sc.total_marker):
                break
            cls = _norm_class(b)
            cl = ClassLoss(size_class=cls)
            for el, (_, tcol) in ecol.items():
                cl.tonnes[el] = _num(grid.get((r, tcol)))
                cl.cells[el] = f"{get_column_letter(tcol)}{r}"
            classes[cls] = cl
            r += 1

        # --- блоки минералогии по классам (ниже таблицы, до конца листа) ---
        for (row, col), v in sorted(grid.items()):
            if col != sc.anchor_col or row <= start:
                continue
            s = str(v)
            if sc.size_unit_marker not in s or sc.total_marker in s:
                continue
            cls = _norm_class(s)
            if cls not in classes:
                continue
            self._read_mineralogy(grid, row, classes[cls], ecol)

        # порядок крупный→тонкий — динамически по числу в подписи, без белого списка
        prof.classes = sorted(classes.values(), key=lambda c: class_sort_key(c.size_class))
        self._validate(prof)
        return prof

    def _validate(self, prof: TailingsProfile):
        """Страховка: ловим признаки сбитого парса, чтобы не выдавать уверенную чушь."""
        if len(prof.classes) < 4:
            prof.warnings.append(
                f"распознано лишь {len(prof.classes)} классов крупности (ожидалось ~6) — "
                "возможно, съехал формат, колонки, или нужна другая схема.")
        el = self.schema.primary_element()
        for cl in prof.classes:
            forms_t = sum(f.tonnes for f in cl.forms if f.element == el)
            class_t = cl.tonnes.get(el) or 0.0
            # баланс: сумма тонн по формам класса ≈ тоннам класса (±25%)
            if class_t and forms_t and abs(forms_t - class_t) / class_t > 0.25:
                prof.warnings.append(
                    f"класс {cl.size_class}: сумма форм {forms_t:.0f} т ≠ итогу класса "
                    f"{class_t:.0f} т (>25%) — проверьте раскладку колонок/схему.")
            rec = cl.recoverable_tonnes(el)
            if class_t and rec > class_t * 1.02:
                prof.warnings.append(
                    f"класс {cl.size_class}: извлекаемого ({rec:.0f} т) больше всего "
                    f"металла класса ({class_t:.0f} т) — раскладка сбита.")

    def _read_mineralogy(self, grid, header_row, cl: ClassLoss, ecol):
        """Прочитать формы блока: строки от header+1 до маркера конца (схема)."""
        sc = self.schema
        forms_seen = {f.form + f.element for f in cl.forms}  # анти-дубли между секциями
        r = header_row + 1
        while r < header_row + 12:
            b = grid.get((r, sc.anchor_col))
            if b is None:
                r += 1; continue
            name = str(b).strip()
            if any(name.startswith(m) or m in name for m in sc.mineralogy_stop_markers):
                break
            for el, (pcol, tcol) in ecol.items():
                t = _num(grid.get((r, tcol)))
                p = _num(grid.get((r, pcol)))
                if t is None and p is None:
                    continue
                key = name + el
                if key in forms_seen:
                    continue
                cl.forms.append(FormLoss(
                    form=name, element=el, tonnes=t or 0.0, pct=p or 0.0,
                    recoverable=sc.is_recoverable(el, name),
                    cell=f"{get_column_letter(tcol)}{r}"))
                forms_seen.add(key)
            r += 1
